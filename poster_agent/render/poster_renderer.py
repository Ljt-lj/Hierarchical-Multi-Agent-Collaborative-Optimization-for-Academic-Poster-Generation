"""学术海报渲染：可读性优化、小圆角卡片、阴影层次、紧凑填充."""

from __future__ import annotations

import textwrap
from pathlib import Path

import re

from PIL import Image, ImageDraw, ImageFont

from poster_agent.config import PosterConfig
from poster_agent.models.render_report import PosterRenderReport, SectionRenderReport
from poster_agent.models.trees import PosterNode
from poster_agent.render.image_fit import (
    figure_aspect_ratio,
    fit_contain_size,
    is_composite_figure_path,
    is_focus_figure_path,
    is_generated_visual,
    is_paper_figure,
    natural_fit_height,
    prepare_paper_figure,
    smart_fit_image,
)
from poster_agent.render.poster_style import clean_figure_caption, figure_panel_title
from poster_agent.render.text_layout import text_width, truncate_to_width, wrap_to_width
from poster_agent.render.text_utils import format_body

THEME = {
    "header_bg": (32, 58, 110),
    "header_fg": (255, 255, 255),
    "header_sub": (190, 205, 230),
    "canvas_bg": (235, 238, 242),
    "section_bg": (255, 255, 255),
    "section_border": (175, 192, 218),
    "section_bar": (41, 84, 144),
    "body_text": (30, 36, 48),
    "shadow": (195, 202, 214),
}

ACADEMIC_THEME = {
    "header_bg": (255, 255, 255),
    "header_fg": (25, 25, 25),
    "header_sub": (70, 70, 70),
    "canvas_bg": (250, 250, 250),
    "section_bg": (255, 255, 255),
    "section_border": (220, 220, 220),
    "section_bar": (196, 92, 38),
    "body_text": (28, 28, 28),
    "shadow": (235, 235, 235),
}

CARD_RADIUS = 4
BAR_RADIUS = 4
SECTION_INNER_PAD = 20
LINE_SPACING = 10
PARA_SPACING = 12
SHADOW_OFFSET = (2, 3)


class _FitTracker:
    """跟踪单区块 bullets 实际渲染数量."""

    def __init__(self) -> None:
        self.max_next_bullet = 0

    def update(self, next_bullet: int) -> None:
        self.max_next_bullet = max(self.max_next_bullet, next_bullet)


class PosterRenderer:
    def __init__(self, output_dir: Path, *, academic_style: bool = True, poster_config: PosterConfig | None = None):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.poster_config = poster_config or PosterConfig()
        self.academic_style = academic_style
        self.theme = ACADEMIC_THEME if academic_style else THEME
        self.CANVAS_W = self.poster_config.width
        self.CANVAS_H = self.poster_config.height
        self.paper_title = "Academic Poster"
        self.authors = ""
        self.last_render_report = PosterRenderReport()

    def render_png(
        self,
        tree: PosterNode,
        name: str,
        paper_title: str | None = None,
        authors: str | None = None,
    ) -> Path:
        if paper_title:
            self.paper_title = paper_title
        if authors:
            self.authors = format_body(authors)
        img = Image.new("RGB", (self.CANVAS_W, self.CANVAS_H), self.theme["canvas_bg"])
        draw = ImageDraw.Draw(img)
        self.last_render_report = PosterRenderReport()
        self._draw_header(draw)
        for node in _collect_leaves(tree):
            self._draw_section(draw, img, node)
        self._draw_footer(draw)
        out = self.output_dir / f"{name}.png"
        img.save(out, quality=96)
        return out

    def render_pptx(
        self,
        tree: PosterNode,
        name: str,
        paper_title: str | None = None,
        authors: str | None = None,
    ) -> Path:
        from pptx import Presentation

        png = self.render_png(tree, f"{name}_tmp", paper_title, authors)
        prs = Presentation()
        prs.slide_width = 9144000 * self.CANVAS_W // 1000
        prs.slide_height = 9144000 * self.CANVAS_H // 1000
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        png.unlink(missing_ok=True)
        out = self.output_dir / f"{name}.pptx"
        prs.save(out)
        return out

    def _draw_header(self, draw: ImageDraw.ImageDraw) -> None:
        landscape = self.CANVAS_W > self.CANVAS_H
        h = 200 if (landscape and self.academic_style) else 280
        th = self.theme
        draw.rectangle([0, 0, self.CANVAS_W, h], fill=th["header_bg"])
        if not self.academic_style:
            draw.rectangle([0, h - 4, self.CANVAS_W, h], fill=(55, 90, 150))

        title_size = 58 if self.academic_style else 64
        sub_size = 24 if self.academic_style else 26
        ft = _load_font(title_size, bold=True)
        fa = _load_font(sub_size)

        title = self.paper_title
        max_title_w = self.CANVAS_W - 200
        title_lines = wrap_to_width(title, ft, max_title_w)
        if len(title_lines) > 2:
            title_lines = title_lines[:2]
            title_lines[-1] = truncate_to_width(title_lines[-1], ft, max_title_w)

        sub = self.authors.strip()
        if not sub and not self.academic_style:
            sub = "Auto-generated Academic Poster | Multi-agent System"
        sub_lines = wrap_to_width(sub, fa, max_title_w) if sub else []
        sub_line = sub_lines[0] if sub_lines else ""
        if len(sub_lines) > 1:
            sub_line = truncate_to_width(sub_line, fa, max_title_w)

        title_line_h = title_size + 12
        sub_line_h = sub_size + 10 if sub_line else 0
        block_h = len(title_lines) * title_line_h + (14 + sub_line_h if sub_line else 0)
        start_y = max(28, (h - block_h) // 2)

        y = start_y
        for line in title_lines:
            tw = text_width(line, ft)
            draw.text(((self.CANVAS_W - tw) // 2, y), line, fill=th["header_fg"], font=ft)
            y += title_line_h

        if sub_line:
            sw = text_width(sub_line, fa)
            draw.text(((self.CANVAS_W - sw) // 2, y + 10), sub_line, fill=th["header_sub"], font=fa)

        if self.academic_style:
            draw.line([(60, h - 2), (self.CANVAS_W - 60, h - 2)], fill=(200, 200, 200), width=2)
        else:
            draw.rounded_rectangle([40, 40, 116, 116], radius=4, outline=(120, 150, 200), width=2)
            draw.text((58, 68), "LOGO", fill=th["header_sub"], font=_load_font(22))

    def _draw_footer(self, draw: ImageDraw.ImageDraw) -> None:
        if self.academic_style:
            return
        y = self.CANVAS_H - 46
        draw.rectangle([0, y, self.CANVAS_W, self.CANVAS_H], fill=THEME["header_bg"])
        msg = "Generated by Hierarchical Multi-agent Poster System"
        fw = _text_width(msg, _load_font(18))
        draw.text(((self.CANVAS_W - fw) // 2, y + 12), msg, fill=THEME["header_sub"], font=_load_font(18))

    def _draw_section(self, draw: ImageDraw.ImageDraw, img: Image.Image, node: PosterNode) -> None:
        r = node.rect
        ox, oy, ow, oh = r.x, r.y, r.width, r.height
        th = self.theme

        if not self.academic_style:
            sx, sy = SHADOW_OFFSET
            draw.rounded_rectangle(
                [ox + sx, oy + sy, ox + ow + sx, oy + oh + sy],
                radius=CARD_RADIUS,
                fill=th["shadow"],
            )
        draw.rounded_rectangle(
            [ox, oy, ox + ow, oy + oh],
            radius=CARD_RADIUS if not self.academic_style else 2,
            fill=th["section_bg"],
            outline=th["section_border"],
            width=1 if self.academic_style else 2,
        )

        title_font = node.font_size_title
        bar_x0 = ox + SECTION_INNER_PAD
        bar_x1 = ox + ow - SECTION_INNER_PAD
        bar_y1 = oy + SECTION_INNER_PAD

        if self.academic_style:
            bar_h = max(title_font + 18, 44)
            draw.rectangle(
                [bar_x0, bar_y1, bar_x1, bar_y1 + bar_h],
                fill=th["section_bar"],
            )
            ft = _load_font(title_font, bold=True)
            section_title = node.title
            bar_text_w = bar_x1 - bar_x0 - 24
            text_y = bar_y1 + (bar_h - title_font) // 2
            draw.text(
                (bar_x0 + 12, text_y),
                truncate_to_width(section_title, ft, bar_text_w),
                fill=(255, 255, 255),
                font=ft,
            )
            content_top = bar_y1 + bar_h + 12
        else:
            bar_h = max(title_font + 24, 48)
            draw.rounded_rectangle(
                [bar_x0, bar_y1, bar_x1, bar_y1 + bar_h],
                radius=BAR_RADIUS,
                fill=th["section_bar"],
            )
            ft = _load_font(title_font, bold=True)
            section_title = node.title
            icon_size = max(int(title_font * 0.65), 22)
            bar_text_w = bar_x1 - bar_x0 - icon_size - 36
            while text_width(section_title, ft) > bar_text_w and title_font > 24:
                title_font -= 2
                ft = _load_font(title_font, bold=True)
            icon_y = bar_y1 + (bar_h - icon_size) // 2
            _draw_bar_icon(draw, bar_x0 + 12, icon_y, node.section_icon or "doc", icon_size)
            text_y = bar_y1 + (bar_h - title_font) // 2
            draw.text(
                (bar_x0 + icon_size + 22, text_y),
                truncate_to_width(section_title, ft, bar_text_w),
                fill=(255, 255, 255),
                font=ft,
            )
            content_top = bar_y1 + bar_h + SECTION_INNER_PAD
        content_box = (
            bar_x0,
            content_top,
            bar_x1,
            oy + oh - SECTION_INNER_PAD,
        )
        mode = node.layout_mode or "text_only"
        figures = _collect_figures(node)
        fb = _load_font(node.font_size_body)
        fit = _FitTracker()

        if mode == "figure_wrap" and figures:
            self._layout_figure_wrap(draw, img, node, content_box, figures, fb, fit)
        elif mode == "figure_rich" and figures:
            self._layout_figure_rich(draw, img, node, content_box, figures, fb, fit)
        elif mode == "figure_top" and figures:
            self._layout_figure_top(draw, img, node, content_box, figures, fb, fit)
        elif mode in ("figure_adaptive", "side_by_side") and figures:
            self._layout_figure_adaptive(draw, img, node, content_box, figures, fb, fit)
        elif mode == "figure_bottom" and figures:
            ratio = 0.58 if _has_paper_figure(figures) else 0.50
            vtype = getattr(node, "visual_type", "") or ""
            if vtype == "logic_pipeline":
                ratio = 0.68
            elif vtype == "bullet_cards":
                ratio = 0.58
            self._layout_figure_bottom(draw, img, node, content_box, figures, fb, fit, fig_ratio=ratio)
        elif mode == "text_dense":
            self._layout_text_dense(draw, node, content_box, fb, fit)
        elif mode == "panel_stack":
            self._layout_panel_stack(draw, img, node, content_box, fb, fit)
        elif mode == "figure_panel":
            self._layout_figure_panel(draw, img, node, content_box, figures, fb, fit)
        else:
            self._layout_text_only(draw, img, node, content_box, figures, fb, fit)

        bullets_total = len(node.bullets)
        bullets_shown = min(fit.max_next_bullet, bullets_total)
        has_visual = bool(node.visual_path and Path(node.visual_path).exists())
        paper_paths = [p for p in node.image_paths if p and Path(p).exists()]
        has_paper = bool(paper_paths) and not (
            has_visual and len(paper_paths) == 1 and paper_paths[0] == node.visual_path
        )
        char_count = len(node.summary or "") + sum(len(b) for b in node.bullets[:bullets_shown])
        self.last_render_report.sections.append(
            SectionRenderReport(
                title=node.title,
                bullets_total=bullets_total,
                bullets_shown=bullets_shown,
                has_visual=has_visual,
                has_paper_figure=has_paper or bool(paper_paths),
                body_font=node.font_size_body,
                area=node.rect.area,
                char_count=char_count,
            )
        )

    def _layout_figure_adaptive(self, draw, img, node, box, figures, fb, fit: _FitTracker) -> None:
        if _has_paper_figure(figures):
            self._layout_figure_wrap(draw, img, node, box, figures, fb, fit)
            return
        aspect = figure_aspect_ratio(figures[0])
        if aspect >= 1.08:
            self._layout_figure_top(draw, img, node, box, figures, fb, fit)
        else:
            self._layout_text_wrap(draw, img, node, box, figures, fb, side="right", fit=fit)

    def _layout_figure_wrap(
        self,
        draw,
        img,
        node,
        box,
        figures,
        fb,
        fit: _FitTracker,
    ) -> None:
        """论文插图：完整显示 + 文字环绕（侧栏 + 下方通栏续排）."""
        x0, y0, x1, y1 = box
        w, h = x1 - x0, y1 - y0
        gap = 12
        paper_figs = [f for f in figures if is_paper_figure(f)]
        extra_figs = [f for f in figures if f not in paper_figs]
        main_fig = paper_figs[0] if paper_figs else figures[0]
        # 始终侧栏环绕：宽图也缩小贴侧，避免顶栏大图裁切/占满
        side = "right"
        fig_frac = 0.36
        fig_w = max(int(w * fig_frac), int(w * 0.30))
        max_fig_h = int(h * 0.52)
        nat_h = self._natural_figure_height(main_fig, fig_w, paper=True)
        fig_h = min(nat_h, max_fig_h)
        fig_h = max(fig_h, int(w * 0.14))

        if side == "right":
            fig_x = x1 - fig_w
            text_x = x0
            text_w_side = max(fig_x - x0 - gap, int(w * 0.34))
        else:
            fig_x = x0
            text_x = fig_x + fig_w + gap
            text_w_side = max(x1 - text_x, int(w * 0.34))

        cy = y0
        next_bullet = 0
        summary_done = False
        if node.summary and text_w_side >= int(w * 0.30):
            cy, next_bullet = self._draw_text_block(
                draw,
                node,
                x0,
                y0,
                w,
                min(int(h * 0.22), fig_h),
                fb,
                numbered=False,
                bullet_start=0,
                include_summary=True,
                stop_between_bullets=False,
                fit=fit,
            )
            summary_done = True
            cy += gap // 2

        fig_y = cy if cy > y0 else y0
        fig_bottom = self._paste_figures(
            img,
            [main_fig],
            fig_x,
            fig_y,
            fig_w,
            fig_h,
            draw,
            fit_contain=True,
            paper_contain=is_paper_figure(main_fig),
        )
        cap = _caption_for_figure(node, main_fig, index=0)
        if cap:
            fig_bottom = _draw_figure_caption(draw, cap, fig_x, fig_bottom, fig_w, fb)
        text_col_h = max(fig_bottom - fig_y, y1 - fig_y - gap)
        cy, next_bullet = self._draw_text_block(
            draw,
            node,
            text_x,
            fig_y,
            text_w_side,
            text_col_h,
            fb,
            numbered=True,
            bullet_start=next_bullet,
            include_summary=not summary_done,
            stop_between_bullets=True,
            fit=fit,
        )
        flow_y = max(fig_bottom, cy) + gap
        remain_h = max(y1 - flow_y, 36)
        final_y = flow_y
        if next_bullet < len(node.bullets) or remain_h > 50:
            final_y, _ = self._draw_text_block(
                draw,
                node,
                x0,
                flow_y,
                w,
                remain_h,
                fb,
                numbered=True,
                bullet_start=next_bullet,
                include_summary=False,
                distribute=True,
                fit=fit,
            )
        if extra_figs and final_y < y1 - 60:
            tbl_y = final_y + gap
            tbl_h = min(int(w * 0.24), max(y1 - tbl_y - 8, 72))
            if tbl_h >= 60:
                tbl_bottom = self._paste_figures(
                    img, extra_figs[:1], x0, tbl_y, w, tbl_h, draw, fit_contain=True,
                )
                if node.visual_type == "data_table":
                    tbl_cap = node.table_caption or "Table: key benchmark metrics"
                    _draw_figure_caption(
                        draw, tbl_cap, x0, tbl_bottom, w, fb, font_size=max(fb.size - 6, 20),
                    )

    def _layout_figure_top_contain(
        self, draw, img, node, box, figures, fb, fit: _FitTracker,
    ) -> None:
        """论文 Fig 通栏完整显示（支持 1–2 张纵向堆叠），正文压缩在下方."""
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        paper_figs = [f for f in figures if is_paper_figure(f)][:2]
        if not paper_figs:
            paper_figs = figures[:2]

        fig_w = int(w * 0.99)
        fig_x = x0 + (w - fig_w) // 2
        cap_reserve = 48
        text_reserve = min(int(total_h * 0.10), 110)
        fp = paper_figs[0]
        nat_h = self._natural_figure_height(fp, fig_w, paper=True)
        fig_budget = min(max(int(total_h * 0.72), nat_h + cap_reserve), total_h - text_reserve - 36)
        gap = 10

        if len(paper_figs) >= 2:
            slot_h = max((fig_budget - gap) // 2, int(fig_w * 0.24))
            cy = y0
            for fp in paper_figs[:2]:
                cap_idx = figures.index(fp) if fp in figures else 0
                nat_h = self._natural_figure_height(fp, fig_w, paper=True)
                use_h = min(nat_h, slot_h)
                bottom = self._paste_figures(
                    img, [fp], fig_x, cy, fig_w, use_h, draw,
                    fit_contain=True, paper_contain=True,
                )
                cap = _caption_for_figure(node, fp, index=cap_idx)
                if cap:
                    bottom = _draw_figure_caption(draw, cap, fig_x, bottom, fig_w, fb)
                cy = bottom + gap
            fig_bottom = cy
        else:
            fp = paper_figs[0]
            fig_bottom = self._paste_figures(
                img, [fp], fig_x, y0, fig_w, fig_budget, draw,
                fit_contain=True, paper_contain=True,
            )
            cap = _caption_for_figure(node, fp, index=0)
            if cap:
                fig_bottom = _draw_figure_caption(draw, cap, fig_x, fig_bottom, fig_w, fb)

        text_y = fig_bottom + 6
        text_h = max(y1 - text_y, 36)
        self._draw_text_block(
            draw, node, x0, text_y, w, text_h, fb,
            numbered=True, compact=True, distribute=False, fit=fit,
        )

    def _layout_figure_rich(
        self, draw, img, node, box, figures, fb, fit: _FitTracker,
    ) -> None:
        """混排：论文 Fig 优先完整显示 + 合成图表行 + 精简正文."""
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        gap = 6

        gen_figs = [f for f in figures if is_generated_visual(f)]
        paper_figs = [f for f in figures if is_paper_figure(f) and not is_generated_visual(f)]

        text_slot = max(int(total_h * 0.11), 72)
        gen_slot = int(total_h * 0.13) if gen_figs else 0
        paper_slot = max(total_h - text_slot - gen_slot - gap * (2 if gen_figs else 1), int(total_h * 0.55))

        cy = y0
        if paper_figs:
            fp = paper_figs[0]
            fig_w = int(w * 0.99)
            fig_x = x0 + (w - fig_w) // 2
            self._paste_figures(
                img, [fp], fig_x, cy, fig_w, paper_slot, draw,
                fit_contain=True, paper_contain=True,
            )
            cap = _caption_for_figure(node, fp, index=0)
            if cap:
                _draw_figure_caption(draw, cap, fig_x, y0 + paper_slot - 40, fig_w, fb)
            cy = y0 + paper_slot + gap

        if gen_figs:
            if len(gen_figs) >= 2:
                half_w = (w - gap) // 2
                cy2 = self._paste_figures(
                    img, [gen_figs[1]], x0 + half_w + gap, cy, half_w, gen_slot, draw,
                    fit_contain=False, fill=True,
                )
                cy = self._paste_figures(
                    img, [gen_figs[0]], x0, cy, half_w, gen_slot, draw,
                    fit_contain=False, fill=True,
                )
                cy = max(cy, cy2) + gap
            else:
                cy = self._paste_figures(
                    img, [gen_figs[0]], x0, cy, w, gen_slot, draw,
                    fit_contain=False, fill=True,
                ) + gap
            cy = max(cy, y0 + paper_slot + gen_slot + gap * 2)

        text_y = cy
        text_h = max(y1 - text_y, 32)
        self._draw_text_block(
            draw, node, x0, text_y, w, text_h, fb,
            numbered=True, compact=True, distribute=False, fit=fit,
        )

    def _layout_text_wrap(
        self,
        draw,
        img,
        node,
        box,
        figures,
        fb,
        *,
        side: str = "right",
        fit: _FitTracker | None = None,
    ) -> None:
        """瘦高插图置侧，正文环绕（先并排、后通栏续排）."""
        x0, y0, x1, y1 = box
        w, h = x1 - x0, y1 - y0
        gap = 12
        fig_w = int(w * 0.40)
        fig_h = min(
            self._natural_figure_height(figures[0], fig_w),
            int(h * 0.58),
        )
        fig_h = max(fig_h, int(w * 0.16))

        if side == "right":
            fig_x = x1 - fig_w
            text_x = x0
            text_w = max(fig_x - x0 - gap, int(w * 0.38))
        else:
            fig_x = x0
            text_x = fig_x + fig_w + gap
            text_w = max(x1 - text_x, int(w * 0.38))

        fig_bottom = self._paste_figures(
            img, figures, fig_x, y0, fig_w, fig_h, draw,
            fit_contain=True, paper_contain=_has_paper_figure(figures),
        )
        cy, next_bullet = self._draw_text_block(
            draw,
            node,
            text_x,
            y0,
            text_w,
            fig_h,
            fb,
            numbered=True,
            bullet_start=0,
            include_summary=True,
            stop_between_bullets=True,
            fit=fit,
        )
        flow_y = max(fig_bottom, cy) + gap
        remain_h = max(y1 - flow_y, 36)
        if next_bullet < len(node.bullets):
            _, next_bullet = self._draw_text_block(
                draw,
                node,
                x0,
                flow_y,
                w,
                remain_h,
                fb,
                numbered=True,
                bullet_start=next_bullet,
                include_summary=False,
                distribute=True,
                fit=fit,
            )
        elif remain_h > 50:
            _, next_bullet = self._draw_text_block(
                draw,
                node,
                x0,
                flow_y,
                w,
                remain_h,
                fb,
                numbered=True,
                bullet_start=len(node.bullets),
                include_summary=False,
                distribute=True,
                fit=fit,
            )

    def _natural_figure_height(self, path: str, width: int, *, paper: bool = False) -> int:
        try:
            src = Image.open(path)
            gen = is_generated_visual(path)
            is_paper = paper or is_paper_figure(path)
            return natural_fit_height(src, width, is_generated=gen, is_paper=is_paper) + 12
        except Exception:
            return int(width * 0.28)

    def _layout_figure_top(self, draw, img, node, box, figures, fb, fit: _FitTracker) -> None:
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        paper = _has_paper_figure(figures)
        if paper:
            self._layout_figure_top_contain(draw, img, node, box, figures, fb, fit)
            return
        vis_type = getattr(node, "visual_type", "") or ""
        max_ratio = 0.68 if vis_type in ("architecture", "bar_chart", "logic_pipeline") else 0.55
        fig_h = self._figure_slot_height(figures, w, total_h, max_ratio=max_ratio)
        fig_bottom = self._paste_figures(
            img, figures, x0, y0, w, fig_h, draw, fit_contain=False, fill=True,
        )
        text_y = fig_bottom + 8
        text_h = max(y1 - text_y, 40)
        self._draw_text_block(
            draw, node, x0, text_y, w, text_h, fb,
            numbered=True, distribute=True, fit=fit,
        )

    def _layout_figure_bottom(self, draw, img, node, box, figures, fb, fit: _FitTracker, *, fig_ratio: float = 0.42) -> None:
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        paper = _has_paper_figure(figures)
        if paper:
            if figures and is_composite_figure_path(figures[0]):
                self._layout_figure_top_contain(draw, img, node, box, figures, fb, fit)
            else:
                self._layout_figure_wrap(draw, img, node, box, figures, fb, fit)
            return
        fig_h = self._figure_slot_height(figures, w, total_h, max_ratio=fig_ratio)
        vtype = getattr(node, "visual_type", "") or ""
        measured = self._measure_text_height(node, w, fb, numbered=True) + 20
        text_h = min(max(measured, int(total_h * 0.20)), int(total_h * 0.42))
        text_bottom, _ = self._draw_text_block(
            draw, node, x0, y0, w, text_h, fb,
            numbered=True, compact=True, distribute=True, fit=fit,
        )
        fig_y = text_bottom + 6
        fig_alloc = max(y1 - fig_y - 4, int(total_h * (0.55 if vtype == "logic_pipeline" else 0.50)))
        gen = figures and is_generated_visual(figures[0])
        self._paste_figures(
            img, figures, x0, fig_y, w, fig_alloc, draw,
            fit_contain=not gen, fill=True,
        )

    def _layout_panel_stack(self, draw, img, node, box, fb, fit: _FitTracker) -> None:
        """Results/Methods：橙色大标题 + 多个 panel（每 panel 上方有小标题）."""
        x0, y0, x1, y1 = box
        w = x1 - x0
        cy = y0
        gap = 10
        panels = node.children or []
        if not panels:
            self._layout_text_dense(draw, node, box, fb, fit)
            return
        n = len(panels)
        gap_total = gap * max(n - 1, 0)
        avail = max(y1 - y0 - gap_total, 100)
        title_l = node.title.lower()
        if title_l == "results" and n >= 4:
            fracs = [0.28, 0.28, 0.24, 0.20][:n]
            s = sum(fracs)
            fracs = [f / s for f in fracs]
            heights = [int(avail * f) for f in fracs]
            if sum(heights) > avail:
                scale = avail / max(sum(heights), 1)
                heights = [max(int(h * scale), 52) for h in heights]
            diff = avail - sum(heights)
            heights[-1] += diff
        elif title_l == "methods" and n == 2:
            heights = [max(int(avail * 0.38), 80), max(avail - int(avail * 0.38), 100)]
        else:
            weights = [max(p.rect.height, 80) for p in panels]
            tw = sum(weights) or float(n)
            heights = [max(int(avail * w / tw), 72) for w in weights[:-1]]
            heights.append(max(avail - sum(heights), 72))
        for panel, ph in zip(panels, heights):
            if cy >= y1 - 28:
                ph = max(y1 - cy, 40)
            if ph <= 0:
                break
            if cy + ph > y1:
                ph = max(y1 - cy, 40)
            panel_box = (x0, cy, x1, cy + ph)
            self._layout_figure_panel(
                draw, img, panel, panel_box,
                _collect_figures(panel), fb, fit,
            )
            cy += ph + gap

    def _layout_figure_panel(
        self, draw, img, node, box, figures, fb, fit: _FitTracker,
    ) -> None:
        """Panel：粗体小标题在插图上方，Fig 编号在插图下方."""
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        cy = y0
        title_fs = max(node.font_size_body + 2, 28)
        title_font = _load_font(title_fs, bold=True)
        title_lines = wrap_to_width(node.title, title_font, w - 8)[:2]
        for line in title_lines:
            draw.text((x0 + 2, cy), line, fill=(28, 28, 28), font=title_font)
            cy += title_fs + 4
        cy += 4

        text_reserve = max(int(total_h * 0.14), 48) if not figures else max(int(total_h * 0.12), 40)
        if paper_paths := [f for f in (figures or []) if is_paper_figure(f) and not is_generated_visual(f)]:
            text_reserve = max(int(total_h * 0.06), 28) if not node.bullets else max(int(total_h * 0.10), 36)
        fig_h = max(total_h - (cy - y0) - text_reserve, int(w * (0.38 if paper_paths else 0.20)))

        if figures:
            cap_raw = _caption_for_figure(node, figures[0], index=0)
            fig_label = clean_figure_caption(cap_raw, figures[0])
            m = re.match(r"^(Fig\.?\s*\d+[a-z]?)", fig_label, re.I)
            fig_ref = m.group(1) if m else ""
            gen_paths = [f for f in figures if is_generated_visual(f)]
            paper_paths = [f for f in figures if is_paper_figure(f) and not is_generated_visual(f)]
            if len(gen_paths) >= 2:
                half_w = (w - 8) // 2
                row_h = max(fig_h, int(w * 0.16))
                cy2 = self._paste_figures(
                    img, [gen_paths[1]], x0 + half_w + 8, cy, half_w, row_h, draw,
                    fit_contain=False, fill=True,
                )
                cy = self._paste_figures(
                    img, [gen_paths[0]], x0, cy, half_w, row_h, draw,
                    fit_contain=False, fill=True,
                )
                cy = max(cy, cy2)
            elif len(gen_paths) == 1:
                use_h = max(fig_h, int(w * (0.16 if node.bullets else 0.22)))
                cy = self._paste_figures(
                    img, gen_paths[:1], x0, cy, w, use_h, draw,
                    fit_contain=False, fill=True,
                )
            else:
                use_h = max(fig_h, int(w * (0.42 if paper_paths else 0.28)))
                cy = self._paste_figures(
                    img, paper_paths[:1] if paper_paths else figures[:1],
                    x0, cy, w, use_h, draw,
                    fit_contain=True, paper_contain=bool(paper_paths),
                )
            if fig_ref and paper_paths:
                cy = _draw_figure_caption(draw, fig_ref, x0, cy, w, fb, font_size=18)
            cy += 4

        text_h = max(y1 - cy, 36)
        self._draw_text_block(
            draw, node, x0, cy, w, text_h, fb,
            numbered=True, compact=True, distribute=False, fit=fit,
        )

    def _layout_text_dense(self, draw, node, box, fb, fit: _FitTracker) -> None:
        x0, y0, x1, y1 = box
        self._draw_text_block(
            draw, node, x0, y0, x1 - x0, y1 - y0, fb,
            numbered=True, compact=True, distribute=False, fit=fit,
        )

    def _layout_text_only(self, draw, img, node, box, figures, fb, fit: _FitTracker) -> None:
        x0, y0, x1, y1 = box
        w, total_h = x1 - x0, y1 - y0
        if figures:
            paper = _has_paper_figure(figures)
            if paper:
                self._layout_figure_wrap(draw, img, node, box, figures, fb, fit)
                return
            fig_h = self._figure_slot_height(figures, w, total_h, max_ratio=0.40)
            text_h = total_h - fig_h - 10
            self._draw_text_block(
                draw, node, x0, y0, w, text_h, fb,
                numbered=True, distribute=True, fit=fit,
            )
            self._paste_figures(img, figures, x0, y0 + text_h + 10, w, fig_h, draw, fill=True)
        else:
            self._draw_text_block(
                draw, node, x0, y0, w, total_h, fb,
                numbered=True, distribute=False, fit=fit,
            )

    def _figure_slot_height(self, figures: list[str], width: int, total_h: int, *, max_ratio: float) -> int:
        if not figures or width <= 0:
            return 0
        pad = 12
        inner_w = max(width - pad * 2, 40)
        n = min(len(figures), 2)
        sub_w = (inner_w - 10) // 2 if n >= 2 else inner_w
        natural = 0
        for fig in figures[:n]:
            try:
                from PIL import Image
                src = Image.open(fig)
                gen = is_generated_visual(fig)
                paper = is_paper_figure(fig)
                nh = natural_fit_height(src, sub_w, is_generated=gen, is_paper=paper)
                if gen and Path(fig).stem.lower().find("vis") >= 0:
                    nh = int(nh * 1.12)
                natural = max(natural, nh + pad * 2)
            except Exception:
                natural = max(natural, int(width * 0.22))
        cap = int(total_h * max_ratio)
        floor_h = int(width * 0.18 if any(is_paper_figure(f) for f in figures[:n]) else width * 0.20)
        return max(floor_h, min(natural, cap, total_h - 80))

    def _measure_text_height(
        self,
        node: PosterNode,
        w: int,
        font,
        *,
        numbered: bool = False,
    ) -> int:
        line_h = font.size + LINE_SPACING
        lines = 0
        if node.summary:
            lines += len(wrap_to_width(format_body(node.summary), font, w))
        for i, bullet in enumerate(node.bullets):
            prefix = f"{i + 1}. " if numbered else ""
            lines += len(wrap_to_width(prefix + format_body(bullet), font, w - 16))
        return lines * line_h + (PARA_SPACING if node.summary and node.bullets else 0)

    def _draw_text_block(
        self,
        draw: ImageDraw.ImageDraw,
        node: PosterNode,
        x: int,
        y: int,
        w: int,
        h: int,
        font,
        *,
        numbered: bool = False,
        compact: bool = False,
        distribute: bool = False,
        fill_vertical: bool = False,
        bullet_start: int = 0,
        include_summary: bool = True,
        stop_between_bullets: bool = False,
        fit: _FitTracker | None = None,
    ) -> tuple[int, int]:
        lines: list[tuple[str, int, int | None]] = []
        summary = format_body(node.summary) if include_summary else ""
        if summary:
            for line in wrap_to_width(summary, font, w):
                lines.append((line, 0, None))
            lines.append(("", 0, None))

        for i, bullet in enumerate(node.bullets):
            if i < bullet_start:
                continue
            text = format_body(bullet)
            prefix = f"{i + 1}. " if numbered else ""
            wrapped = wrap_to_width(prefix + text, font, w - 16)
            for j, line in enumerate(wrapped):
                lines.append((line, 16 if j > 0 else 0, i))

        if not lines:
            return y, bullet_start

        base_spacing = 6 if compact else LINE_SPACING
        content_lines = sum(1 for line, _, _ in lines if line)

        def _block_height(extra_sp: int, body_font) -> int:
            lh = body_font.size + base_spacing + extra_sp
            total = 0
            for line, _, _ in lines:
                if line:
                    total += lh
                else:
                    total += PARA_SPACING // 2
            return total

        extra_spacing = 0
        if distribute and h > 0 and len(lines) > 1:
            if fill_vertical and h > 280 and content_lines >= 2:
                for _ in range(4):
                    max_extra = max(
                        72,
                        int((h - PARA_SPACING * 2) / max(content_lines, 1) - font.size - base_spacing),
                    )
                    lo, hi, best, best_fill = 0, max(max_extra, 0), 0, 0
                    while lo <= hi:
                        mid = (lo + hi) // 2
                        bh = _block_height(mid, font)
                        if bh <= h:
                            if bh >= best_fill:
                                best = mid
                                best_fill = bh
                            lo = mid + 1
                        else:
                            hi = mid - 1
                    extra_spacing = best
                    if best_fill >= h * 0.88 or font.size >= 52:
                        break
                    font = _load_font(min(font.size + 3, 52))
            elif distribute:
                extra = max(0, h - _block_height(0, font))
                cap = 22 if (extra > 80 and content_lines < 14) else (16 if extra > 40 else 10)
                extra_spacing = min(extra // max(len(lines) - 1, 1), cap)
        line_h = font.size + base_spacing + extra_spacing

        cy = y
        bottom = y + h
        next_bullet = bullet_start
        for line, indent, bullet_idx in lines:
            if not line:
                cy += PARA_SPACING // 2
                continue
            if cy + line_h > bottom:
                if stop_between_bullets and bullet_idx is not None:
                    return cy, bullet_idx
                break
            draw.text((x + indent, cy), line, fill=THEME["body_text"], font=font)
            cy += line_h
            if bullet_idx is not None:
                next_bullet = bullet_idx + 1
        if fit is not None:
            fit.update(next_bullet)
        return cy, next_bullet

    def _paste_figures(
        self,
        img: Image.Image,
        figures: list[str],
        x: int,
        y: int,
        w: int,
        h: int,
        draw: ImageDraw.ImageDraw,
        *,
        fill: bool = False,
        fit_contain: bool = False,
        paper_contain: bool = False,
    ) -> int:
        if not figures or w <= 8 or h <= 8:
            return y
        pad = 6
        inner_x0 = x + pad
        inner_y0 = y + pad
        inner_w = w - pad * 2
        inner_h = h - pad * 2
        if inner_w <= 8 or inner_h <= 8:
            return y

        n = min(len(figures), 2)
        gap = 10
        use_h = inner_h
        if paper_contain or fit_contain:
            use_h = self._paper_contain_height(figures[0], inner_w, inner_h)
        elif not fill:
            slot_h = self._figure_slot_height(figures, w, h + 120, max_ratio=1.0)
            use_h = min(inner_h, max(slot_h - pad * 2, int(inner_w * 0.22)))
        if fill:
            use_h = inner_h

        use_fill = fill and not fit_contain and not paper_contain
        th = self.theme
        border = th["section_border"]
        is_paper = paper_contain or (fit_contain and figures and is_paper_figure(figures[0]))
        if n == 1:
            if is_paper:
                draw.rectangle(
                    [inner_x0 - 1, inner_y0 - 1, inner_x0 + inner_w + 1, inner_y0 + use_h + 1],
                    fill=(252, 252, 252),
                    outline=border,
                    width=1,
                )
            used_h = self._fit_image(
                img, figures[0], inner_x0, inner_y0, inner_w, use_h,
                fill=use_fill,
                paper_contain=is_paper,
            )
            if used_h > 0:
                use_h = used_h
        else:
            sub_w = (inner_w - gap) // 2
            if not use_fill:
                draw.rounded_rectangle(
                    [inner_x0, inner_y0, inner_x0 + inner_w, inner_y0 + use_h],
                    radius=CARD_RADIUS,
                    outline=THEME["section_border"],
                    width=1,
                )
            tight = use_fill or use_h >= inner_h * 0.85
            self._fit_image(img, figures[0], inner_x0, inner_y0, sub_w, use_h, fill=tight and not fit_contain)
            self._fit_image(
                img, figures[1], inner_x0 + sub_w + gap, inner_y0, sub_w, use_h, fill=tight and not fit_contain
            )
        if n == 1 and not is_paper and not use_fill:
            draw.rectangle(
                [inner_x0, inner_y0, inner_x0 + inner_w, inner_y0 + use_h],
                outline=border,
                width=1,
            )
        return inner_y0 + use_h + pad

    def _paper_contain_height(self, path: str, max_w: int, max_h: int) -> int:
        try:
            src = prepare_paper_figure(Image.open(path).convert("RGB"), path=path)
            _, nh = fit_contain_size(src.size[0], src.size[1], max_w, max_h)
            return max(nh, int(max_w * 0.22))
        except Exception:
            return int(max_w * 0.32)

    def _fit_image(
        self,
        canvas: Image.Image,
        path: str,
        x: int,
        y: int,
        w: int,
        h: int,
        *,
        fill: bool = False,
        paper_contain: bool = False,
    ) -> int:
        """粘贴插图；paper_contain 时返回实际占用高度."""
        try:
            if not Path(path).exists() or w <= 0 or h <= 0:
                return 0
            src = Image.open(path).convert("RGB")
            paper = is_paper_figure(path)
            gen = is_generated_visual(path)
            if paper_contain or (paper and not fill):
                src = prepare_paper_figure(src, path=path)
                nw, nh = fit_contain_size(src.size[0], src.size[1], w, h)
                fitted = src.resize((nw, nh), Image.Resampling.LANCZOS)
                px = x + (w - nw) // 2
                draw = ImageDraw.Draw(canvas)
                draw.rounded_rectangle(
                    [px - 2, y - 2, px + nw + 2, y + nh + 2],
                    radius=CARD_RADIUS,
                    outline=THEME["section_border"],
                    width=1,
                )
                canvas.paste(fitted, (px, y))
                return nh
            if fill and gen:
                sw, sh = src.size
                if sw > 0 and sh > 0:
                    nh = h
                    nw = min(max(1, int(sw * nh / sh)), w)
                    fitted = src.resize((nw, nh), Image.Resampling.LANCZOS)
                    px = x + (w - nw) // 2
                    canvas.paste(fitted, (px, y))
                    return h
            fitted = smart_fit_image(src, w, h, is_generated=gen, is_paper=paper)
            canvas.paste(fitted, (x, y))
            return h if fill else 0
        except Exception as exc:
            import logging
            logging.getLogger(__name__).warning("fit_image failed %s: %s", path, exc)
            return 0


def _fit_contain(img: Image.Image, tw: int, th: int) -> Image.Image:
    """等比缩放至区域内完整显示，浅底居中，不裁剪."""
    sw, sh = img.size
    scale = min(tw / sw, th / sh)
    nw, nh = max(1, int(sw * scale)), max(1, int(sh * scale))
    img = img.resize((nw, nh), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (tw, th), (252, 253, 255))
    canvas.paste(img, ((tw - nw) // 2, (th - nh) // 2))
    return canvas


def _cover_resize(img: Image.Image, tw: int, th: int) -> Image.Image:
    """等比放大并居中裁剪，填满区域无留白."""
    sw, sh = img.size
    scale = max(tw / sw, th / sh)
    nw, nh = int(sw * scale), int(sh * scale)
    img = img.resize((nw, nh), Image.Resampling.LANCZOS)
    left = (nw - tw) // 2
    top = (nh - th) // 2
    return img.crop((left, top, left + tw, top + th))


def _draw_bar_icon(draw: ImageDraw.ImageDraw, x: int, y: int, icon: str, size: int = 24) -> None:
    s = size
    draw.rounded_rectangle([x, y, x + s, y + s], radius=3, fill=(70, 110, 170))
    cx, cy = x + s // 2, y + s // 2
    scale = s / 24
    white = (255, 255, 255)
    if icon == "chart":
        draw.rectangle([cx - int(6 * scale), cy, cx - int(3 * scale), cy + int(5 * scale)], fill=white)
        draw.rectangle([cx - int(1 * scale), cy - int(4 * scale), cx + int(2 * scale), cy + int(5 * scale)], fill=white)
        draw.rectangle([cx + int(4 * scale), cy + int(1 * scale), cx + int(7 * scale), cy + int(5 * scale)], fill=white)
    elif icon == "gear":
        draw.ellipse([cx - int(6 * scale), cy - int(6 * scale), cx + int(6 * scale), cy + int(6 * scale)], outline=white, width=max(1, int(scale)))
        draw.ellipse([cx - int(2 * scale), cy - int(2 * scale), cx + int(2 * scale), cy + int(2 * scale)], fill=white)
    elif icon == "lightbulb":
        draw.ellipse([cx - int(5 * scale), cy - int(7 * scale), cx + int(5 * scale), cy + int(3 * scale)], fill=white)
        draw.rectangle([cx - int(3 * scale), cy + int(3 * scale), cx + int(3 * scale), cy + int(7 * scale)], fill=white)
    elif icon == "flag":
        draw.line([(cx - int(4 * scale), cy - int(7 * scale)), (cx - int(4 * scale), cy + int(7 * scale))], fill=white, width=max(1, int(2 * scale)))
        draw.polygon([
            (cx - int(4 * scale), cy - int(7 * scale)),
            (cx + int(7 * scale), cy - int(3 * scale)),
            (cx - int(4 * scale), cy + int(1 * scale)),
        ], fill=white)
    else:
        draw.rectangle([cx - int(6 * scale), cy - int(5 * scale), cx + int(6 * scale), cy + int(6 * scale)], outline=white, width=max(1, int(scale)))


def _has_paper_figure(figures: list[str]) -> bool:
    return any(is_paper_figure(f) for f in figures)


def _collect_leaves(node: PosterNode) -> list[PosterNode]:
    if node.title in ("Poster Root", "Section Group"):
        return [c for ch in node.children for c in _collect_leaves(ch)]
    if node.children and (
        node.layout_mode == "panel_stack"
        or any(c.block_style == "panel" for c in node.children)
    ):
        return [node]
    if not node.children:
        return [node]
    return [c for ch in node.children for c in _collect_leaves(ch)]


def _collect_figures(node: PosterNode) -> list[str]:
    """收集插图：论文 Fig 优先，其次合成图."""
    paths: list[str] = []
    for p in node.image_paths:
        if not p or not Path(p).exists() or p in paths:
            continue
        paths.append(p)
    if node.visual_path and Path(node.visual_path).exists() and node.visual_path not in paths:
        paths.append(node.visual_path)
    elif node.image_path and Path(node.image_path).exists() and node.image_path not in paths:
        paths.append(node.image_path)
    paper = [p for p in paths if is_paper_figure(p)]
    other = [p for p in paths if p not in paper]
    return (paper + other)[:4]


def _load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    names = ["msyhbd.ttc", "msyh.ttc", "simhei.ttf", "arialbd.ttf" if bold else "arial.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _text_width(text: str, font) -> int:
    try:
        return int(font.getlength(text))
    except Exception:
        return len(text) * (font.size if hasattr(font, "size") else 12)


def _wrap_lines(text: str, width: int) -> list[str]:
    return textwrap.wrap(text, width=width) or [text]


def _caption_for_figure(node: PosterNode, path: str, *, index: int = 0) -> str:
    raw = ""
    if index < len(node.figure_captions):
        raw = node.figure_captions[index]
    return clean_figure_caption(raw, path)


def _table_caption_label(node: PosterNode) -> str:
    return node.table_caption or ""


def _draw_figure_caption(
    draw: ImageDraw.ImageDraw,
    caption: str,
    x: int,
    y: int,
    w: int,
    body_font,
    *,
    font_size: int | None = None,
) -> int:
    if not caption.strip():
        return y
    fs = font_size or max(body_font.size - 8, 16)
    font = _load_font(fs)
    lines = wrap_to_width(caption, font, w - 8)[:2]
    cy = y + 6
    for line in lines:
        m = re.match(r"^(Fig\.?\s*\d+[a-z]?)\s*:?\s*(.*)", line, re.I)
        if m:
            fig_label, rest = m.group(1), m.group(2)
            draw.text((x + 4, cy), fig_label + (": " if rest else ""), fill=(196, 92, 38), font=_load_font(fs, bold=True))
            lw = _text_width(fig_label + ": ", _load_font(fs, bold=True))
            if rest:
                draw.text((x + 4 + lw, cy), rest, fill=(85, 85, 85), font=font)
        else:
            draw.text((x + 4, cy), line, fill=(85, 85, 85), font=font)
        cy += fs + 3
    return cy + 2
